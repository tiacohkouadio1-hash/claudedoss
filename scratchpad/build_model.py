#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Construit le modèle financier PV (Côte d'Ivoire) — modele-financier-pv.xlsx"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/claudedoss/docs/modele-financier-pv.xlsx"

# ---- Styles ----
ARIAL = "Arial"
def font(color="FF000000", bold=False, size=10, italic=False):
    return Font(name=ARIAL, size=size, bold=bold, color=color, italic=italic)

BLUE  = "FF0000FF"   # input hardcodé
BLACK = "FF000000"   # formule
GREEN = "FF008000"   # lien vers une autre feuille

HDR_FILL   = PatternFill("solid", fgColor="1F4E5C")
SUB_FILL   = PatternFill("solid", fgColor="DCE6EA")
YEL_FILL   = PatternFill("solid", fgColor="FFFF00")
TITLE_FILL = PatternFill("solid", fgColor="F5A623")

thin = Side(style="thin", color="B7C4CC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Formats
F_FCFA = '#,##0" FCFA";(#,##0)" FCFA";"-"'
F_KWH  = '#,##0" kWh"'
F_KWC  = '#,##0.0" kWc"'
F_PCT  = '0.0%'
F_TAR  = '#,##0" FCFA/kWh"'
F_YRS  = '#,##0.0" ans"'
F_M2   = '#,##0" m²"'
F_INT  = '#,##0'

wb = Workbook()

def put(ws, ref, value, color=BLACK, bold=False, fmt=None, fill=None,
        align=None, size=10, border=False, italic=False):
    c = ws[ref]
    c.value = value
    c.font = font(color=color, bold=bold, size=size, italic=italic)
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(align=="wrap"))
    if border: c.border = BORDER
    return c

def title(ws, text, span="A1:D1"):
    a = span.split(":")[0]
    put(ws, a, text, color="FF14212B", bold=True, size=14, fill=TITLE_FILL, align="left")
    ws.merge_cells(span)
    for row in ws[span]:
        for c in row:
            c.fill = TITLE_FILL

# =====================================================================
# 1) HYPOTHÈSES
# =====================================================================
h = wb.active
h.title = "Hypothèses"
h.sheet_view.showGridLines = False
title(h, "Modèle financier PV — Hypothèses", "A1:D1")

put(h, "A2", "Légende : cellules en BLEU = à modifier (hypothèses) · JAUNE = paramètres clés · NOIR = calculées automatiquement.",
    color="FF5B6B78", italic=True, size=9)
h.merge_cells("A2:D2")

# --- Paramètres généraux ---
put(h, "A4", "Paramètres généraux", color="FFFFFFFF", bold=True, fill=HDR_FILL)
for ref in ("B4","C4"):
    put(h, ref, "", fill=HDR_FILL)
put(h, "A4", "Paramètre", color="FFFFFFFF", bold=True, fill=HDR_FILL)
put(h, "B4", "Valeur", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")
put(h, "C4", "Unité", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")

gen = [
    ("Productible spécifique (Abidjan)", 1400, "kWh/kWc/an", BLUE, YEL_FILL),
    ("Dégradation annuelle des modules", 0.007, "/an", BLUE, None),
    ("Puissance par panneau", 500, "Wc", BLUE, None),
    ("Surface par kWc", 6.5, "m²/kWc", BLUE, None),
    ("Tarif électricité initial", 80, "FCFA/kWh", BLUE, YEL_FILL),
    ("Inflation du tarif électricité", 0.03, "/an", BLUE, None),
    ("O&M annuel (% du CAPEX)", 0.015, "% CAPEX", BLUE, None),
    ("Durée d'analyse", 25, "ans", BLUE, None),
    ("Taux d'actualisation (VAN)", 0.08, "/an", BLUE, None),
]
r = 5
for lbl, val, unit, col, fill in gen:
    put(h, f"A{r}", lbl, border=True)
    fmt = F_PCT if isinstance(val, float) and val < 1 else F_INT
    put(h, f"B{r}", val, color=col, fmt=fmt, fill=fill, align="center", border=True)
    put(h, f"C{r}", unit, color="FF5B6B78", align="center", border=True)
    r += 1
# B5 productible, B6 degr, B7 wc, B8 m2, B9 tarif, B10 infl, B11 om, B12 duree, B13 actu

# --- Paramètres par configuration ---
put(h, "A15", "Paramètres par configuration", color="FFFFFFFF", bold=True, fill=HDR_FILL)
put(h, "B15", "B1 · Injection directe", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")
put(h, "C15", "B2 · Hybride (batterie)", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")
put(h, "D15", "B3 · Sur-mesure", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")

cfg = [
    ("Coût unitaire (FCFA/kWc)", 700000, 1500000, 1100000, F_INT, YEL_FILL),
    ("Taux d'autoconsommation", 0.65, 0.90, 0.80, F_PCT, None),
    ("Remplacement onduleur an 12 (% CAPEX)", 0.08, 0.08, 0.08, F_PCT, None),
    ("Remplacement batterie an 10 (% CAPEX)", 0.00, 0.30, 0.15, F_PCT, None),
]
r = 16
for lbl, b1, b2, b3, fmt, fill in cfg:
    put(h, f"A{r}", lbl, border=True)
    for col, v in zip("BCD", (b1, b2, b3)):
        put(h, f"{col}{r}", v, color=BLUE, fmt=fmt, fill=fill, align="center", border=True)
    r += 1
# Row 16 coût, 17 taux, 18 onduleur, 19 batterie

put(h, "A21", "Sources : voir docs/etude-photovoltaique-abidjan.md et docs/business-plan.md. "
              "Coûts et tarif à confirmer par devis fournisseurs / factures CIE.",
    color="FF5B6B78", italic=True, size=9)
h.merge_cells("A21:D21")

h.column_dimensions["A"].width = 40
for col in "BCD":
    h.column_dimensions[col].width = 22

# Références utiles (adresses)
P_PRODUCTIBLE = "Hypothèses!$B$5"
P_DEGR        = "Hypothèses!$B$6"
P_WC          = "Hypothèses!$B$7"
P_M2          = "Hypothèses!$B$8"
P_TARIF0      = "Hypothèses!$B$9"
P_INFL        = "Hypothèses!$B$10"
P_OM          = "Hypothèses!$B$11"
P_DUREE       = "Hypothèses!$B$12"
P_ACTU        = "Hypothèses!$B$13"
CFG_COST = {"B1":"Hypothèses!$B$16","B2":"Hypothèses!$C$16","B3":"Hypothèses!$D$16"}
CFG_AUTO = {"B1":"Hypothèses!$B$17","B2":"Hypothèses!$C$17","B3":"Hypothèses!$D$17"}
CFG_OND  = {"B1":"Hypothèses!$B$18","B2":"Hypothèses!$C$18","B3":"Hypothèses!$D$18"}
CFG_BAT  = {"B1":"Hypothèses!$B$19","B2":"Hypothèses!$C$19","B3":"Hypothèses!$D$19"}

# =====================================================================
# 2) CAS CLIENT
# =====================================================================
cc = wb.create_sheet("Cas client")
cc.sheet_view.showGridLines = False
title(cc, "Cas client — dimensionnement à partir de la facture", "A1:C1")

put(cc, "A2", "Renseignez les 3 cellules jaunes ci-dessous ; le reste se calcule tout seul.",
    color="FF5B6B78", italic=True, size=9)
cc.merge_cells("A2:C2")

put(cc, "A4", "Données client (à remplir)", color="FFFFFFFF", bold=True, fill=HDR_FILL)
put(cc, "B4", "Valeur", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")
put(cc, "C4", "Unité", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center")

put(cc, "A5", "Facture mensuelle moyenne", border=True)
put(cc, "B5", 45000, color=BLUE, fmt=F_FCFA, fill=YEL_FILL, align="center", border=True)
put(cc, "C5", "FCFA/mois", color="FF5B6B78", align="center", border=True)

put(cc, "A6", "Tarif électricité appliqué", border=True)
put(cc, "B6", f"={P_TARIF0}", color=GREEN, fmt=F_TAR, fill=YEL_FILL, align="center", border=True)
put(cc, "C6", "FCFA/kWh", color="FF5B6B78", align="center", border=True)

put(cc, "A7", "Taux de couverture solaire visé", border=True)
put(cc, "B7", 0.80, color=BLUE, fmt=F_PCT, fill=YEL_FILL, align="center", border=True)
put(cc, "C7", "% conso", color="FF5B6B78", align="center", border=True)

put(cc, "A9", "Dimensionnement (calculé)", color="FFFFFFFF", bold=True, fill=HDR_FILL)
put(cc, "B9", "", fill=HDR_FILL); put(cc, "C9", "", fill=HDR_FILL)

put(cc, "A10", "Consommation mensuelle", border=True)
put(cc, "B10", "=B5/B6", fmt=F_KWH, align="center", border=True)
put(cc, "C10", "kWh/mois", color="FF5B6B78", align="center", border=True)

put(cc, "A11", "Consommation annuelle", border=True)
put(cc, "B11", "=B10*12", fmt=F_KWH, align="center", border=True, bold=True)
put(cc, "C11", "kWh/an", color="FF5B6B78", align="center", border=True)

put(cc, "A12", "Consommation à couvrir", border=True)
put(cc, "B12", "=B11*B7", fmt=F_KWH, align="center", border=True)
put(cc, "C12", "kWh/an", color="FF5B6B78", align="center", border=True)

put(cc, "A13", "Puissance crête recommandée", border=True)
put(cc, "B13", f"=B12/{P_PRODUCTIBLE}", fmt=F_KWC, align="center", border=True, bold=True)
put(cc, "C13", "kWc", color="FF5B6B78", align="center", border=True)

put(cc, "A14", "Nombre de panneaux", border=True)
put(cc, "B14", f"=ROUNDUP(B13*1000/{P_WC},0)", fmt=F_INT, align="center", border=True)
put(cc, "C14", "panneaux", color="FF5B6B78", align="center", border=True)

put(cc, "A15", "Surface de toiture estimée", border=True)
put(cc, "B15", f"=B13*{P_M2}", fmt=F_M2, align="center", border=True)
put(cc, "C15", "m²", color="FF5B6B78", align="center", border=True)

cc.column_dimensions["A"].width = 34
cc.column_dimensions["B"].width = 18
cc.column_dimensions["C"].width = 14

CLIENT_KWC   = "'Cas client'!$B$13"
CLIENT_CONSO = "'Cas client'!$B$11"
CLIENT_TARIF = "'Cas client'!$B$6"

# =====================================================================
# 3) CASH-FLOW (une feuille par config)
# =====================================================================
def build_cashflow(code, name):
    ws = wb.create_sheet(f"Cash-flow {code}")
    ws.sheet_view.showGridLines = False
    title(ws, f"Cash-flow 25 ans — {name}", "A1:I1")

    # Bloc paramètres (liens)
    put(ws, "A3", "Puissance", bold=True); put(ws, "B3", f"={CLIENT_KWC}", color=GREEN, fmt=F_KWC)
    put(ws, "A4", "CAPEX", bold=True);     put(ws, "B4", f"=B3*{CFG_COST[code]}", fmt=F_FCFA, bold=True)
    put(ws, "D3", "Conso annuelle", bold=True); put(ws, "E3", f"={CLIENT_CONSO}", color=GREEN, fmt=F_KWH)
    put(ws, "D4", "Taux autoconso.", bold=True); put(ws, "E4", f"={CFG_AUTO[code]}", color=GREEN, fmt=F_PCT)

    # Indicateurs
    put(ws, "G3", "VAN (25 ans)", bold=True)
    put(ws, "H3", f"=NPV({P_ACTU},H14:H38)+H13", color="FF14212B", bold=True, fmt=F_FCFA)
    put(ws, "G4", "Année de rentabilité", bold=True)
    # 1re année où le flux cumulé devient positif (robuste même si le cumul
    # rechute l'année d'un remplacement batterie) via la colonne d'aide J.
    put(ws, "H4", f"=IF(MIN(J13:J38)>{P_DUREE},\">25 ans\",MIN(J13:J38))", bold=True, fmt=F_YRS)

    # En-têtes du tableau
    hdrs = ["Année","Production (kWh)","Tarif (FCFA/kWh)","Autoconso. (kWh)",
            "Économie brute","O&M","Remplacements","Flux net","Flux cumulé"]
    for i, txt in enumerate(hdrs):
        col = get_column_letter(1+i)
        put(ws, f"{col}12", txt, color="FFFFFFFF", bold=True, fill=HDR_FILL,
            align="center", border=True, size=9)

    CAPEX = "$B$4"; KWC = "$B$3"; CONSO = "$E$3"; TAUX = "$E$4"
    # Année 0
    row = 13
    put(ws, f"A{row}", 0, align="center", border=True, fmt=F_INT)
    for col in "BCDEF": put(ws, f"{col}{row}", 0, fmt=F_KWH if col in "BD" else F_FCFA, align="center", border=True)
    put(ws, f"G{row}", 0, fmt=F_FCFA, align="center", border=True)
    put(ws, f"H{row}", f"=-{CAPEX}", fmt=F_FCFA, align="center", border=True)
    put(ws, f"I{row}", f"=H{row}", fmt=F_FCFA, align="center", border=True)
    # Colonne d'aide J (masquée) : année si cumulé ≥ 0, sinon grande valeur
    put(ws, f"J{row}", f"=IF(I{row}>=0,A{row},9999)", fmt=F_INT)
    # Années 1..25
    for n in range(1, 26):
        row = 13 + n
        pr = row - 1
        put(ws, f"A{row}", n, align="center", border=True, fmt=F_INT)
        put(ws, f"B{row}", f"={KWC}*{P_PRODUCTIBLE}*(1-{P_DEGR})^(A{row}-1)", fmt=F_INT, align="center", border=True)
        put(ws, f"C{row}", f"={P_TARIF0}*(1+{P_INFL})^(A{row}-1)", fmt=F_INT, align="center", border=True)
        put(ws, f"D{row}", f"=MIN(B{row}*{TAUX},{CONSO})", fmt=F_INT, align="center", border=True)
        put(ws, f"E{row}", f"=D{row}*C{row}", fmt=F_FCFA, align="center", border=True)
        put(ws, f"F{row}", f"={CAPEX}*{P_OM}", fmt=F_FCFA, align="center", border=True)
        put(ws, f"G{row}", f"=IF(A{row}=10,{CAPEX}*{CFG_BAT[code]},0)+IF(A{row}=12,{CAPEX}*{CFG_OND[code]},0)",
            fmt=F_FCFA, align="center", border=True)
        put(ws, f"H{row}", f"=E{row}-F{row}-G{row}", fmt=F_FCFA, align="center", border=True)
        put(ws, f"I{row}", f"=I{pr}+H{row}", fmt=F_FCFA, align="center", border=True)
        put(ws, f"J{row}", f"=IF(I{row}>=0,A{row},9999)", fmt=F_INT)

    ws.column_dimensions["A"].width = 8
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["J"].hidden = True  # colonne d'aide
    return ws

build_cashflow("B1", "B1 · Injection directe (sans batterie)")
build_cashflow("B2", "B2 · Hybride (avec batterie)")
build_cashflow("B3", "B3 · Sur-mesure")

# =====================================================================
# 4) COMPARATIF CONFIGS
# =====================================================================
cp = wb.create_sheet("Comparatif configs")
cp.sheet_view.showGridLines = False
title(cp, "Comparatif des 3 configurations", "A1:D1")
put(cp, "A2", "Même puissance (depuis « Cas client ») comparée sur les 3 architectures.",
    color="FF5B6B78", italic=True, size=9)
cp.merge_cells("A2:D2")

put(cp, "A4", "Indicateur", color="FFFFFFFF", bold=True, fill=HDR_FILL, border=True)
put(cp, "B4", "B1 · Injection", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center", border=True)
put(cp, "C4", "B2 · Hybride", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center", border=True)
put(cp, "D4", "B3 · Sur-mesure", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center", border=True)

cols = {"B":"B1","C":"B2","D":"B3"}
CFsheet = {"B1":"'Cash-flow B1'","B2":"'Cash-flow B2'","B3":"'Cash-flow B3'"}

rows = [
    ("Puissance (kWc)", lambda c: f"={CLIENT_KWC}", F_KWC, False),
    ("CAPEX (investissement)", lambda c: f"={CLIENT_KWC}*{CFG_COST[cols[c]]}", F_FCFA, True),
    ("Production année 1", lambda c: f"={CLIENT_KWC}*{P_PRODUCTIBLE}", F_INT, False),
    ("Énergie autoconsommée an 1", lambda c: f"=MIN({CLIENT_KWC}*{P_PRODUCTIBLE}*{CFG_AUTO[cols[c]]},{CLIENT_CONSO})", F_INT, False),
    ("Économie brute an 1", lambda c: f"=MIN({CLIENT_KWC}*{P_PRODUCTIBLE}*{CFG_AUTO[cols[c]]},{CLIENT_CONSO})*{CLIENT_TARIF}", F_FCFA, False),
    ("O&M an 1", lambda c: f"={CLIENT_KWC}*{CFG_COST[cols[c]]}*{P_OM}", F_FCFA, False),
    ("Économie nette an 1", lambda c: f"={c}9-{c}10", F_FCFA, True),
    ("Payback brut (années)", lambda c: f"=IFERROR({c}6/{c}11,\"n/a\")", F_YRS, True),
    ("VAN 25 ans (actualisée)", lambda c: f"={CFsheet[cols[c]]}!$H$3", F_FCFA, True),
    ("Flux cumulé à 25 ans", lambda c: f"={CFsheet[cols[c]]}!$I$38", F_FCFA, False),
]
r = 5
for lbl, fn, fmt, bold in rows:
    put(cp, f"A{r}", lbl, border=True, bold=bold)
    for c in "BCD":
        color = GREEN if lbl.startswith("VAN") or lbl.startswith("Flux cumulé") or lbl=="Puissance (kWc)" else BLACK
        put(cp, f"{c}{r}", fn(c), color=color, fmt=fmt, align="center", border=True, bold=bold)
    r += 1

put(cp, "A16", "Lecture : B1 = meilleur payback (pas de batterie). B2 = confort/nuit (clim) mais retour plus long. "
               "VAN positive = projet créateur de valeur sur 25 ans.",
    color="FF5B6B78", italic=True, size=9)
cp.merge_cells("A16:D16")

cp.column_dimensions["A"].width = 30
for c in "BCD":
    cp.column_dimensions[c].width = 20

# =====================================================================
# 5) SENSIBILITÉ (config B1)
# =====================================================================
sn = wb.create_sheet("Sensibilité")
sn.sheet_view.showGridLines = False
title(sn, "Sensibilité du payback — Config B1 (injection directe)", "A1:F1")
put(sn, "A2", "Payback brut (années) selon le tarif électricité (lignes) et le coût du kWc (colonnes).",
    color="FF5B6B78", italic=True, size=9)
sn.merge_cells("A2:F2")

# En-tête colonnes = coûts unitaires
costs = [600000, 700000, 800000, 900000]
put(sn, "A4", "Tarif \\ Coût", color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center", border=True)
for j, cost in enumerate(costs):
    col = get_column_letter(2+j)
    put(sn, f"{col}4", cost, color="FFFFFFFF", bold=True, fill=HDR_FILL, align="center", border=True, fmt=F_INT)

tarifs = [60, 70, 80, 90, 100, 110, 120]
AUTO_B1 = CFG_AUTO["B1"]
for i, tar in enumerate(tarifs):
    row = 5 + i
    put(sn, f"A{row}", tar, color=BLUE, bold=True, fill=SUB_FILL, align="center", border=True, fmt=F_TAR)
    for j, cost in enumerate(costs):
        col = get_column_letter(2+j)
        colL = f"{col}$4"        # coût (header colonne)
        rowT = f"$A{row}"        # tarif (header ligne)
        capex = f"({CLIENT_KWC}*{colL})"
        prod  = f"({CLIENT_KWC}*{P_PRODUCTIBLE})"
        auto  = f"MIN({prod}*{AUTO_B1},{CLIENT_CONSO})"
        eco   = f"({auto}*{rowT})"
        om    = f"({capex}*{P_OM})"
        f = f"=IFERROR({capex}/({eco}-{om}),\"n/a\")"
        put(sn, f"{col}{row}", f, fmt=F_YRS, align="center", border=True)

put(sn, "A13", "Colonnes = coût unitaire (FCFA/kWc) · Lignes = tarif évité (FCFA/kWh). "
               "Plus le tarif monte / le coût baisse, plus le retour est rapide.",
    color="FF5B6B78", italic=True, size=9)
sn.merge_cells("A13:F13")

sn.column_dimensions["A"].width = 16
for j in range(len(costs)):
    sn.column_dimensions[get_column_letter(2+j)].width = 16

# ---- Ordre des onglets ----
wb.move_sheet("Cas client", -(wb.sheetnames.index("Cas client")-1))
order = ["Hypothèses","Cas client","Comparatif configs","Cash-flow B1","Cash-flow B2","Cash-flow B3","Sensibilité"]
wb._sheets.sort(key=lambda s: order.index(s.title))

# Forcer le recalcul à l'ouverture (Excel / LibreOffice) puisque openpyxl
# n'écrit pas les valeurs en cache.
wb.calculation.fullCalcOnLoad = True

wb.save(OUT)
print("Écrit :", OUT)
print("Onglets :", wb.sheetnames)
